"""Excel and PDF output.

Excel is the working artefact -- the commercial team will edit it. PDF is the
read-only pack that goes to the client. Both are generated from the same
`Estimate` object so they can never disagree.
"""

from __future__ import annotations

from pathlib import Path

from .schema import Estimate

# --------------------------------------------------------------------------
# Shared palette
# --------------------------------------------------------------------------
NAVY = "1F3A5F"
TEAL = "2E7D8F"
AMBER = "E8A33D"
LIGHT = "EEF2F6"
GREY = "8A94A0"
RED = "C0392B"


# ==========================================================================
# Excel
# ==========================================================================

def to_excel(est: Estimate, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    path = Path(path)
    wb = Workbook()

    hdr_fill = PatternFill("solid", fgColor=NAVY)
    sub_fill = PatternFill("solid", fgColor=TEAL)
    alt_fill = PatternFill("solid", fgColor=LIGHT)
    bar_fill = PatternFill("solid", fgColor=TEAL)
    ms_fill = PatternFill("solid", fgColor=AMBER)
    hdr_font = Font(color="FFFFFF", bold=True, size=11)
    title_font = Font(bold=True, size=14, color=NAVY)
    thin = Side(style="thin", color="D0D6DD")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def header_row(ws, row: int, values: list[str], fill=hdr_fill):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.fill = fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = box

    def widths(ws, spec: dict[str, int]):
        for col, w in spec.items():
            ws.column_dimensions[col].width = w

    # ---------------- Summary ----------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Effort Estimate — {est.brief.project_name}"
    ws["A1"].font = title_font
    ws["A2"] = f"{est.brief.client} · generated {est.generated_at.isoformat()}"
    ws["A2"].font = Font(italic=True, color=GREY)

    rows = [
        ("Engagement type", est.brief.engagement_type.value.replace("_", " ").title()),
        ("Purpose", est.brief.purpose.value.replace("_", " ").title()),
        ("Technology profile", est.brief.tech_profile.value.replace("_", " ").title()),
        ("Proposed tools", ", ".join(est.brief.proposed_tools) or "Not specified"),
        ("Delivery model", est.brief.delivery_model.value.title()),
        ("Start date", est.brief.start_date.isoformat()),
        ("", ""),
        ("Adjusted function points", round(est.size.adjusted_function_points, 1)),
        ("Equivalent KSLOC", est.size.equivalent_ksloc),
        ("Delivery rate (h/FP)", est.calibration.pdr_hours_per_fp),
        ("", ""),
        ("Effort — optimistic (h)", est.effort.optimistic_hours),
        ("Effort — likely (h)", est.effort.likely_hours),
        ("Effort — pessimistic (h)", est.effort.pessimistic_hours),
        ("Effort — likely (person-months)", round(est.effort.likely_person_months, 1)),
        ("", ""),
        ("Requested duration (months)", est.schedule.requested_months),
        ("Nominal duration (months)", est.schedule.nominal_months),
        ("Compression ratio", est.schedule.compression_ratio),
        ("Schedule penalty (SCED)", est.schedule.sced_multiplier),
        ("Recommended duration (months)", est.schedule.recommended_months),
        ("Feasible", "YES" if est.schedule.feasible else "NO"),
        ("", ""),
        ("Peak team size (FTE)", est.resources.peak_team_size),
        ("Average team size (FTE)", est.resources.average_team_size),
    ]
    if est.brief.blended_rate_per_hour > 0:
        rows += [("", ""),
                 (f"Blended rate ({est.brief.currency}/h)", est.brief.blended_rate_per_hour),
                 (f"Indicative cost ({est.brief.currency})", est.estimated_cost)]

    r = 4
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=bool(label))
        cell = ws.cell(row=r, column=2, value=value)
        if label == "Feasible" and value == "NO":
            cell.font = Font(bold=True, color=RED)
        r += 1

    ws.cell(row=r + 1, column=1, value="Schedule verdict").font = Font(bold=True)
    vc = ws.cell(row=r + 1, column=2, value=est.schedule.verdict)
    vc.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 4, end_column=7)

    if est.narrative:
        ws.cell(row=r + 6, column=1, value="Commentary").font = Font(bold=True)
        nc = ws.cell(row=r + 6, column=2, value=est.narrative)
        nc.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r + 6, start_column=2, end_row=r + 20, end_column=7)

    widths(ws, {"A": 34, "B": 22, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14})

    # ---------------- Effort by phase ----------------
    ws = wb.create_sheet("Effort by Phase")
    header_row(ws, 1, ["Phase", "Start (mo)", "End (mo)", "Duration (mo)",
                       "Effort (h)", "% of total", "Person-months", "Key deliverables"])
    for i, p in enumerate(est.phases, start=2):
        ws.cell(row=i, column=1, value=p.name)
        ws.cell(row=i, column=2, value=round(p.start_month, 2))
        ws.cell(row=i, column=3, value=round(p.end_month, 2))
        ws.cell(row=i, column=4, value=round(p.end_month - p.start_month, 2))
        ws.cell(row=i, column=5, value=p.effort_hours)
        ws.cell(row=i, column=6, value=p.effort_pct).number_format = "0.0%"
        ws.cell(row=i, column=7, value=round(p.effort_hours / 152, 1))
        ws.cell(row=i, column=8, value=", ".join(p.deliverables))
        if i % 2 == 0:
            for c in range(1, 9):
                ws.cell(row=i, column=c).fill = alt_fill
    total_row = len(est.phases) + 2
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=5,
            value=f"=SUM(E2:E{total_row - 1})").font = Font(bold=True)
    ws.cell(row=total_row, column=7,
            value=f"=SUM(G2:G{total_row - 1})").font = Font(bold=True)
    widths(ws, {"A": 30, "B": 11, "C": 11, "D": 13, "E": 12, "F": 11, "G": 15, "H": 60})
    ws.freeze_panes = "A2"

    # ---------------- Resource loading ----------------
    ws = wb.create_sheet("Resource Loading")
    n = est.resources.months
    header_row(ws, 1, ["Role"] + [f"M{i + 1}" for i in range(n)] + ["Total PM", "Peak FTE"])
    for i, role in enumerate(est.resources.roles, start=2):
        ws.cell(row=i, column=1, value=role.role).font = Font(bold=True)
        for m, fte in enumerate(role.monthly_ftes):
            c = ws.cell(row=i, column=2 + m, value=fte)
            c.number_format = "0.00"
            if fte > 0:
                intensity = min(fte / max(est.resources.peak_team_size, 0.01), 1.0)
                shade = f"{int(255 - 90 * intensity):02X}{int(255 - 50 * intensity):02X}FF"
                c.fill = PatternFill("solid", fgColor=shade)
        ws.cell(row=i, column=2 + n, value=role.total_person_months).font = Font(bold=True)
        ws.cell(row=i, column=3 + n, value=role.peak_fte)
    tr = len(est.resources.roles) + 2
    ws.cell(row=tr, column=1, value="TOTAL FTE").font = Font(bold=True)
    for m, tot in enumerate(est.resources.monthly_total_ftes):
        c = ws.cell(row=tr, column=2 + m, value=tot)
        c.font = Font(bold=True)
        c.number_format = "0.00"
        c.fill = PatternFill("solid", fgColor=LIGHT)
    widths(ws, {"A": 24})
    for m in range(n + 2):
        ws.column_dimensions[get_column_letter(2 + m)].width = 9
    ws.freeze_panes = "B2"

    # staffing curve chart
    try:
        from openpyxl.chart import LineChart, Reference
        chart = LineChart()
        chart.title = "Total FTE by month"
        chart.y_axis.title = "FTE"
        chart.x_axis.title = "Month"
        data = Reference(ws, min_col=2, max_col=1 + n, min_row=tr, max_row=tr)
        chart.add_data(data, from_rows=True, titles_from_data=False)
        chart.height, chart.width = 8, 22
        ws.add_chart(chart, f"A{tr + 3}")
    except Exception:  # noqa: BLE001
        pass

    # ---------------- Milestones ----------------
    ws = wb.create_sheet("Milestones")
    header_row(ws, 1, ["#", "Milestone", "Month offset", "Target date", "Phase", "Gate criteria"])
    for i, ms in enumerate(est.milestones, start=2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=ms.name).font = Font(bold=True)
        ws.cell(row=i, column=3, value=round(ms.month_offset, 2))
        ws.cell(row=i, column=4, value=ms.calendar_date.isoformat())
        ws.cell(row=i, column=5, value=ms.phase)
        ws.cell(row=i, column=6, value=ms.gate_criteria).alignment = Alignment(wrap_text=True)
        if i % 2 == 0:
            for c in range(1, 7):
                ws.cell(row=i, column=c).fill = alt_fill
    widths(ws, {"A": 5, "B": 32, "C": 14, "D": 14, "E": 28, "F": 60})
    ws.freeze_panes = "A2"

    # ---------------- Gantt ----------------
    ws = wb.create_sheet("Gantt")
    ws["A1"] = f"Project Plan — {est.brief.project_name}"
    ws["A1"].font = title_font
    header_row(ws, 3, ["Phase"] + [f"M{i + 1}" for i in range(n)])
    for i, p in enumerate(est.phases, start=4):
        ws.cell(row=i, column=1, value=p.name).font = Font(bold=True)
        for m in range(n):
            overlap = min(p.end_month, m + 1) - max(p.start_month, m)
            cell = ws.cell(row=i, column=2 + m)
            cell.border = box
            if overlap > 0:
                cell.fill = bar_fill
    ms_row = len(est.phases) + 5
    ws.cell(row=ms_row, column=1, value="Milestones").font = Font(bold=True)
    for ms in est.milestones:
        idx = min(int(ms.month_offset), n - 1)
        cell = ws.cell(row=ms_row, column=2 + idx)
        cell.fill = ms_fill
        cell.value = "◆"
        cell.alignment = Alignment(horizontal="center")
    legend = ms_row + 2
    ws.cell(row=legend, column=1, value="Legend:").font = Font(bold=True)
    ws.cell(row=legend, column=2).fill = bar_fill
    ws.cell(row=legend, column=3, value="Phase duration")
    ws.cell(row=legend + 1, column=2).fill = ms_fill
    ws.cell(row=legend + 1, column=3, value="Milestone / gate")
    widths(ws, {"A": 30})
    for m in range(n):
        ws.column_dimensions[get_column_letter(2 + m)].width = 5

    # ---------------- Assumptions, risks, sources ----------------
    ws = wb.create_sheet("Assumptions & Risks")
    ws["A1"] = "Assumptions"
    ws["A1"].font = title_font
    r = 2
    for a in est.assumptions:
        ws.cell(row=r, column=1, value=f"• {a}").alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        r += 1
    r += 2
    ws.cell(row=r, column=1, value="Risks").font = title_font
    r += 1
    for risk in est.risks:
        cell = ws.cell(row=r, column=1, value=f"• {risk}")
        cell.alignment = Alignment(wrap_text=True)
        if risk.startswith("CRITICAL"):
            cell.font = Font(bold=True, color=RED)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        r += 1
    r += 2
    ws.cell(row=r, column=1, value="Benchmark sources").font = title_font
    r += 1
    ws.cell(row=r, column=1, value=est.calibration.pdr_source)
    r += 1
    for c in est.calibration.citations:
        ws.cell(row=r, column=1, value=c)
        r += 1
    widths(ws, {"A": 110})

    # ---------------- Sizing detail ----------------
    ws = wb.create_sheet("Sizing Detail")
    header_row(ws, 1, ["Component", "Value"])
    for i, (k, v) in enumerate(est.size.breakdown.items(), start=2):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    widths(ws, {"A": 34, "B": 16})

    wb.save(path)
    return path


# ==========================================================================
# PDF
# ==========================================================================

def to_pdf(est: Estimate, path: str | Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    path = Path(path)
    navy = colors.HexColor("#" + NAVY)
    teal = colors.HexColor("#" + TEAL)
    amber = colors.HexColor("#" + AMBER)
    light = colors.HexColor("#" + LIGHT)

    doc = SimpleDocTemplate(
        str(path), pagesize=landscape(A4),
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm,
        title=f"Effort Estimate — {est.brief.project_name}",
    )
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], textColor=navy, fontSize=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=teal, fontSize=12, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", parent=ss["BodyText"], fontSize=8.5, leading=11.5, alignment=TA_LEFT)
    small = ParagraphStyle("small", parent=body, fontSize=7.5, textColor=colors.HexColor("#" + GREY))

    def table(data, widths_mm, header=True, font_size=8):
        t = Table(data, colWidths=[w * mm for w in widths_mm], repeatRows=1 if header else 0)
        style = [
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D6DD")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light]),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        if header:
            style += [("BACKGROUND", (0, 0), (-1, 0), navy),
                      ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                      ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
        t.setStyle(TableStyle(style))
        return t

    story = []

    # --- page 1: summary ---
    story.append(Paragraph(f"Effort Estimate — {est.brief.project_name}", h1))
    story.append(Paragraph(
        f"{est.brief.client} &nbsp;·&nbsp; "
        f"{est.brief.engagement_type.value.replace('_', ' ').title()} / "
        f"{est.brief.purpose.value.replace('_', ' ').title()} &nbsp;·&nbsp; "
        f"Generated {est.generated_at.isoformat()}", small))
    story.append(Spacer(1, 6))

    kpi = [
        ["Adjusted size", "Likely effort", "Effort range", "Duration", "Peak team"],
        [f"{est.size.adjusted_function_points:,.0f} FP",
         f"{est.effort.likely_hours:,.0f} h\n({est.effort.likely_person_months:.1f} PM)",
         f"{est.effort.optimistic_hours:,.0f} – {est.effort.pessimistic_hours:,.0f} h",
         f"{est.schedule.recommended_months:.1f} mo\n(nominal {est.schedule.nominal_months:.1f})",
         f"{est.resources.peak_team_size:.1f} FTE"],
    ]
    k = table(kpi, [52, 52, 58, 52, 52])
    k.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"),
                           ("FONTSIZE", (0, 1), (-1, 1), 11),
                           ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                           ("TEXTCOLOR", (0, 1), (-1, 1), navy),
                           ("TOPPADDING", (0, 1), (-1, 1), 7),
                           ("BOTTOMPADDING", (0, 1), (-1, 1), 7)]))
    story.append(k)

    story.append(Paragraph("Schedule assessment", h2))
    verdict_colour = colors.HexColor("#" + RED) if not est.schedule.feasible else navy
    story.append(Paragraph(est.schedule.verdict,
                           ParagraphStyle("v", parent=body, textColor=verdict_colour)))

    if est.narrative:
        story.append(Paragraph("Commentary", h2))
        for para in [p for p in est.narrative.split("\n") if p.strip()]:
            story.append(Paragraph(para.strip(), body))
            story.append(Spacer(1, 3))

    story.append(Paragraph("Basis of estimate", h2))
    story.append(Paragraph(
        f"Sizing: {est.size.method}. "
        f"Delivery rate: {est.calibration.pdr_hours_per_fp:.2f} hours per function point "
        f"({est.calibration.pdr_source}). "
        f"Effort blended from ISBSG delivery-rate and COCOMO II post-architecture paths. "
        f"Schedule tested against the COCOMO II nominal duration "
        f"TDEV = 3.67 × PM^F with a 75% compression floor.", body))

    # --- page 2: phases + gantt ---
    story.append(PageBreak())
    story.append(Paragraph("Phase plan and effort distribution", h1))
    rows = [["Phase", "Start", "End", "Effort (h)", "% ", "PM", "Key deliverables"]]
    for p in est.phases:
        rows.append([p.name, f"M{p.start_month:.1f}", f"M{p.end_month:.1f}",
                     f"{p.effort_hours:,.0f}", f"{p.effort_pct:.0%}",
                     f"{p.effort_hours / 152:.1f}",
                     Paragraph(", ".join(p.deliverables), small)])
    story.append(table(rows, [46, 16, 16, 22, 14, 16, 137]))

    story.append(Paragraph("Gantt", h2))
    n = est.resources.months
    gantt = [[""] + [f"M{i + 1}" for i in range(n)]]
    for p in est.phases:
        gantt.append([p.name] + [""] * n)
    col_w = max(4.5, min(9.0, 200.0 / max(n, 1)))
    g = Table(gantt, colWidths=[50 * mm] + [col_w * mm] * n)
    gstyle = [
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D6DD")),
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 1), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
    ]
    for ri, p in enumerate(est.phases, start=1):
        for m in range(n):
            if min(p.end_month, m + 1) - max(p.start_month, m) > 0:
                gstyle.append(("BACKGROUND", (1 + m, ri), (1 + m, ri), teal))
    for ms in est.milestones:
        idx = min(int(ms.month_offset), n - 1)
        row = next((i for i, p in enumerate(est.phases, start=1) if p.name == ms.phase), 1)
        gstyle.append(("BACKGROUND", (1 + idx, row), (1 + idx, row), amber))
    g.setStyle(TableStyle(gstyle))
    story.append(g)
    story.append(Spacer(1, 4))
    story.append(Paragraph("Teal = phase duration · Amber = milestone or gate", small))

    # --- page 3: resources ---
    story.append(PageBreak())
    story.append(Paragraph("Resource loading plan (FTE by month)", h1))
    rrows = [["Role"] + [f"M{i + 1}" for i in range(n)] + ["Total PM", "Peak"]]
    for role in est.resources.roles:
        rrows.append([role.role] + [f"{f:.1f}" if f > 0.04 else "–" for f in role.monthly_ftes]
                     + [f"{role.total_person_months:.1f}", f"{role.peak_fte:.1f}"])
    rrows.append(["TOTAL FTE"] + [f"{t:.1f}" for t in est.resources.monthly_total_ftes]
                 + [f"{sum(r.total_person_months for r in est.resources.roles):.1f}",
                    f"{est.resources.peak_team_size:.1f}"])
    rc = max(4.5, min(9.0, 170.0 / max(n, 1)))
    rt = Table(rrows, colWidths=[42 * mm] + [rc * mm] * n + [16 * mm, 12 * mm], repeatRows=1)
    rt.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D6DD")),
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, len(rrows) - 1), (-1, len(rrows) - 1), light),
        ("FONTNAME", (0, len(rrows) - 1), (-1, len(rrows) - 1), "Helvetica-Bold"),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F7F9FB")]),
    ]))
    story.append(rt)

    story.append(Paragraph("Milestones", h2))
    mrows = [["Milestone", "Month", "Target date", "Gate criteria"]]
    for ms in est.milestones:
        mrows.append([ms.name, f"M{ms.month_offset:.1f}", ms.calendar_date.isoformat(),
                      Paragraph(ms.gate_criteria, small)])
    story.append(table(mrows, [55, 16, 24, 172]))

    # --- page 4: assumptions and risks ---
    story.append(PageBreak())
    story.append(Paragraph("Assumptions", h1))
    for a in est.assumptions:
        story.append(Paragraph(f"• {a}", body))
    story.append(Paragraph("Risks", h2))
    for risk in est.risks:
        style = body if not risk.startswith("CRITICAL") else ParagraphStyle(
            "risk", parent=body, textColor=colors.HexColor("#" + RED))
        story.append(Paragraph(f"• {risk}", style))
    story.append(Paragraph("Benchmark sources", h2))
    story.append(Paragraph(est.calibration.pdr_source, small))
    for c in est.calibration.citations:
        story.append(Paragraph(c, small))
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        "This estimate is indicative and produced from parametric models calibrated to "
        "public industry benchmarks. It is not a fixed-price commitment. Replace the "
        "benchmark tables with your own delivery actuals before contractual use.", small))

    doc.build(story)
    return path
